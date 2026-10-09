import io
from docx import Document
from pypdf import PdfReader


def extract_pdf(data: bytes) -> list[tuple[int | None, str]]:
    reader = PdfReader(io.BytesIO(data))
    return [(i, page.extract_text() or "") for i, page in enumerate(reader.pages, start=1)]


def extract_docx(data: bytes) -> list[tuple[int | None, str]]:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
    return [(None, "\n\n".join(parts))]
